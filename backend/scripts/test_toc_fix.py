import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Inches, Pt, RGBColor

doc = docx.Document()
p = doc.add_paragraph('Table of Contents')

# Run 1: begin
r1 = p.add_run()
fld1 = OxmlElement('w:fldChar')
fld1.set(qn('w:fldCharType'), 'begin')
r1._r.append(fld1)

# Run 2: instrText
r2 = p.add_run()
instr = OxmlElement('w:instrText')
instr.set(qn('xml:space'), 'preserve')
instr.text = ' TOC \\o "1-3" \\h \\z \\u '
r2._r.append(instr)

# Run 3: separate
r3 = p.add_run()
fld2 = OxmlElement('w:fldChar')
fld2.set(qn('w:fldCharType'), 'separate')
r3._r.append(fld2)

# Run 4: placeholder text
r4 = p.add_run('Right-click to update Table of Contents')
r4.font.italic = True

# Run 5: end
r5 = p.add_run()
fld3 = OxmlElement('w:fldChar')
fld3.set(qn('w:fldCharType'), 'end')
r5._r.append(fld3)

doc.save('test_toc.docx')
print('test_toc.docx created successfully.')
